from django.db.models import Q, OuterRef, Subquery, DateField, Case, When, Value, CharField, Count
from django import forms
from django.urls import reverse
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from openpyxl import Workbook
from django.utils import timezone
from .models import BienDongNhanKhau, HoKhau, NhanKhau, CanHo
from core.modules.account.models import TaiKhoan
from core.forms import HoKhauForm, NhanKhauForm, BienDongForm
from core.decorators import role_required
from core.modules.base.models import LoaiBienDong
from openpyxl.styles import PatternFill, Font, Alignment


def get_user_hokhau_active(user):
    """Trả về queryset hộ khẩu active mà user là chủ hộ (nếu là cư dân), ngược lại trả về None."""
    if hasattr(user, 'vaitro') and getattr(user.vaitro, 'id_vaitro', None) == 2:
        return HoKhau.objects.filter(is_active=True, is_deleted=False, id_chuho=user)
    return None


@login_required(login_url="login")
@role_required([1, 2])
def nhankhau_list(request):
    search = request.GET.get('search', '').strip()
    loai_biendong = request.GET.get('status', '').strip()
    user_hokhau = get_user_hokhau_active(request.user)
    qs = NhanKhau.objects.filter(
        is_deleted=False, id_hokhau__is_deleted=False)
    if user_hokhau is not None:
        qs = qs.filter(id_hokhau__in=user_hokhau)
    if search:
        qs = qs.filter(
            Q(ho_ten__icontains=search) |
            Q(cccd__icontains=search) |
            Q(so_dien_thoai__icontains=search) |
            Q(email__icontains=search) |
            Q(id_hokhau__id_canho__so_can_ho__icontains=search)
        )
    latest_bd_subquery = BienDongNhanKhau.objects.filter(
        id_nhankhau=OuterRef('pk'),
        is_deleted=False
    ).order_by('-updated_at')
    qs = qs.annotate(
        latest_bd_ngay=Subquery(latest_bd_subquery.values(
            'ngay_batdau')[:1], output_field=DateField()),
        latest_bd_loai=Subquery(latest_bd_subquery.values('loai_biendong')[:1])
    )
    if loai_biendong:
        qs = qs.filter(latest_bd_loai=loai_biendong)
    qs = qs.order_by('-latest_bd_ngay')

    choices_dict = dict(LoaiBienDong.choices)

    qs = qs.annotate(
        latest_bd_loai_display=Case(
            *[
                When(latest_bd_loai=k, then=Value(v))
                for k, v in choices_dict.items()
            ],
            default=Value('-'),
            output_field=CharField()
        )
    )

    return render(request, 'resident/nhankhau/list.html', {
        'nhankhau_list': qs,
        'search': search,
        'loai_biendong': loai_biendong,
        'query': search,
        'status': loai_biendong
    })


@login_required(login_url="login")
@role_required([1, 2])
def nhankhau_detail(request, pk):
    nk = get_object_or_404(NhanKhau, pk=pk, is_deleted=False)
    user_hokhau = get_user_hokhau_active(request.user)
    # Nếu là cư dân chỉ cho xem nhân khẩu thuộc hộ khẩu của mình
    if user_hokhau is not None and nk.id_hokhau not in user_hokhau:
        messages.error(request, 'Bạn không có quyền xem nhân khẩu này!')
        return redirect('nhankhau_list')
    biendong_list = BienDongNhanKhau.objects.filter(
        id_nhankhau=nk, is_deleted=False).order_by('-ngay_batdau')
    return render(request, 'resident/nhankhau/detail.html', {'nk': nk, 'biendong_list': biendong_list})


@login_required(login_url="login")
@role_required([1, 2])
def nhankhau_add(request):
    user_hokhau = get_user_hokhau_active(request.user)
    hokhau_avai = user_hokhau if user_hokhau is not None else HoKhau.objects.filter(
        is_active=True, is_deleted=False)
    if request.method == 'POST':
        form = NhanKhauForm(request.POST, hokhau_avai=hokhau_avai)
        if form.is_valid():
            nk = form.save(commit=False)
            nk.updated_by = request.user.username
            nk.save()
            loai_bd = (
                nk.id_hokhau.resident_status if nk.id_hokhau and nk.id_hokhau.resident_status else 'tam_vang')
            BienDongNhanKhau.objects.create(
                loai_biendong=loai_bd,
                ngay_batdau=timezone.now().date(),
                ly_do="Đăng ký mới",
                id_nhankhau=nk,
                updated_by=request.user.username
            )
            messages.success(request, 'Đăng ký nhân khẩu thành công!')
            return redirect('nhankhau_list')
        else:
            # Lấy tất cả lỗi của form (bao gồm cả lỗi trường và lỗi chung)
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(
                        request, f"{form.fields[field].label if field in form.fields else field}: {error}")
            # Nếu có lỗi chung (non_field_errors)
            for error in form.non_field_errors():
                messages.error(request, error)
            return redirect('nhankhau_list')
    else:
        form = NhanKhauForm(hokhau_avai=hokhau_avai)
    return render(request, 'resident/nhankhau/form.html', {'form': form, 'action': 'add'})


@login_required(login_url="login")
@role_required([1, 2])
def nhankhau_edit(request, pk):
    nk = get_object_or_404(NhanKhau, pk=pk, is_deleted=False)
    user_hokhau = get_user_hokhau_active(request.user)
    hokhau_avai = user_hokhau if user_hokhau is not None else HoKhau.objects.filter(
        is_active=True, is_deleted=False)
    if user_hokhau is not None and nk.id_hokhau not in user_hokhau:
        messages.error(request, 'Bạn không có quyền sửa nhân khẩu này!')
        return redirect('nhankhau_list')

    if request.method == 'POST':
        # Nếu trường id_hokhau bị disabled thì sẽ không có trong POST, nên gán lại từ instance
        post = request.POST.copy()
        post['id_hokhau'] = str(nk.id_hokhau.pk) if nk.id_hokhau else ''
        form = NhanKhauForm(post, instance=nk,
                            hokhau_avai=hokhau_avai, disabled_hokhau=True)
        if form.is_valid():
            nk = form.save(commit=False)
            nk.updated_by = request.user.username
            nk.save()
            messages.success(request, 'Cập nhật nhân khẩu thành công!')
            return redirect('nhankhau_list')
        else:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(
                        request, f"{form.fields[field].label if field in form.fields else field}: {error}")
            for error in form.non_field_errors():
                messages.error(request, error)
            return redirect('nhankhau_list')
    else:
        form = NhanKhauForm(
            instance=nk, hokhau_avai=hokhau_avai, disabled_hokhau=True)
    return render(request, 'resident/nhankhau/edit_form.html', {'form': form, 'action': 'edit', 'nk': nk})


@login_required(login_url="login")
@role_required([1, 2])
def nhankhau_delete(request, pk):
    nk = get_object_or_404(NhanKhau, pk=pk, is_deleted=False)
    user_hokhau = get_user_hokhau_active(request.user)
    if request.method != 'POST':
        return render(request, 'resident/nhankhau/confirm_delete.html', {'nk': nk})

    # Nếu là cư dân chỉ cho xóa nhân khẩu thuộc hộ khẩu của mình
    if user_hokhau is not None and nk.id_hokhau not in user_hokhau:
        messages.error(request, 'Bạn không có quyền xóa nhân khẩu này!')
        return redirect('nhankhau_list')

    # Không cho xóa nếu là chủ hộ của hộ khẩu active
    if nk.is_chu_ho and nk.is_active:
        messages.error(
            request, 'Không thể xóa nhân khẩu đang là chủ hộ!')
        return redirect('nhankhau_list')

    # Không cho xóa nếu có nhiều hơn 1 biến động nhân khẩu
    biendong_count = nk.nhan_khau.filter(is_deleted=False).count()
    if biendong_count > 1:
        messages.error(
            request, 'Không thể xóa nhân khẩu có đăng ký cư trú!')
        return redirect('nhankhau_list')

    nk.is_deleted = True
    nk.updated_by = request.user.username
    nk.save()
    messages.success(request, 'Đã xóa nhân khẩu!')
    return redirect('nhankhau_list')


@login_required(login_url="login")
@role_required([1])
def hokhau_toggle_active(request, pk):
    hokhau = get_object_or_404(HoKhau, pk=pk)
    if request.method == 'POST':
        hokhau.is_active = not hokhau.is_active
        hokhau.updated_by = request.user.username
        hokhau.save(update_fields=['is_active', 'updated_at', 'updated_by'])

        # Lấy tất cả nhân khẩu thuộc hộ khẩu này (không xóa)
        nhankhaus = NhanKhau.objects.filter(id_hokhau=hokhau, is_deleted=False)
        now = timezone.now().date()
        if hokhau.is_active:
            # Khi bật lại
            for nk in nhankhaus:
                nk.is_active = True
                nk.updated_by = request.user.username
                nk.save(update_fields=['is_active', 'updated_by'])
                BienDongNhanKhau.objects.create(
                    id_nhankhau=nk,
                    loai_biendong=hokhau.resident_status if hokhau.resident_status else 'tam_vang',
                    ngay_batdau=now,
                    ly_do='Đăng ký lại',
                    updated_by=request.user.username
                )
        else:
            for nk in nhankhaus:
                nk.is_active = False
                nk.updated_by = request.user.username
                nk.save(update_fields=['is_active', 'updated_by'])
                BienDongNhanKhau.objects.create(
                    id_nhankhau=nk,
                    loai_biendong='cham_dut',
                    ngay_batdau=now,
                    ly_do='Hộ khẩu ngừng hoạt động',
                    updated_by=request.user.username
                )
        messages.success(
            request, f'Trạng thái hộ khẩu đã được cập nhật!')
    return redirect('hokhau_list')


@login_required(login_url="login")
@role_required([1])
def hokhau_list(request):
    hokhau_qs = HoKhau.objects.filter(is_deleted=False).select_related(
        'id_canho', 'id_chuho').annotate(
        nhankhau_count=Count('ho_khau', filter=Q(
            ho_khau__is_deleted=False, ho_khau__is_active=True))
    ).order_by('-id_hokhau')

    search_hk = request.GET.get('search_hk', '').strip()
    resident_status = request.GET.get('resident_status', '')
    is_active = request.GET.get('is_active', '')

    if search_hk:
        hokhau_qs = hokhau_qs.filter(
            Q(id_canho__so_can_ho__icontains=search_hk) |
            Q(id_chuho__username__icontains=search_hk)
        )
    if resident_status:
        hokhau_qs = hokhau_qs.filter(resident_status=resident_status)
    if is_active in ['0', '1']:
        hokhau_qs = hokhau_qs.filter(is_active=(is_active == '1'))

    total_count = hokhau_qs.count()
    return render(request, 'resident/hokhau.html', {
        'hokhau_list': hokhau_qs,
        'search_hk': search_hk,
        'resident_status': resident_status,
        'is_active': is_active,
        'total_count': total_count,
    })


@login_required(login_url="login")
@role_required([1])
def hokhau_add(request):
    # Lấy danh sách căn hộ chưa có hộ khẩu active
    canho_with_active_hokhau = HoKhau.objects.filter(
        is_deleted=False, is_active=True, id_canho__isnull=False).values_list('id_canho', flat=True)
    available_canho = CanHo.objects.filter(is_deleted=False).exclude(
        id_canho__in=canho_with_active_hokhau)
    # Lấy danh sách chủ hộ chưa có hộ khẩu active
    chuho_with_active_hokhau = HoKhau.objects.filter(
        is_deleted=False, is_active=True, id_chuho__isnull=False).values_list('id_chuho', flat=True)
    available_chuho = TaiKhoan.objects.filter(vaitro__id_vaitro=2, is_deleted=False, is_active=True).exclude(
        id_taikhoan__in=chuho_with_active_hokhau)

    if request.method == 'POST':
        form = HoKhauForm(
            request.POST, available_canho=available_canho, available_chuho=available_chuho)
        if form.is_valid():
            id_canho = form.cleaned_data.get('id_canho')
            id_chuho = form.cleaned_data.get('id_chuho')
            # Nếu form là ModelForm thì lấy instance, nếu không thì lấy từ cleaned_data
            if hasattr(form, 'instance') and form.instance:
                id_canho = form.instance.id_canho or id_canho
                id_chuho = form.instance.id_chuho or id_chuho
            # Kiểm tra bản ghi đã tồn tại nhưng đang is_active=False
            existing = None
            if id_canho and id_chuho:
                existing = HoKhau.objects.filter(
                    id_canho=id_canho, id_chuho=id_chuho, is_deleted=False, is_active=False).first()
            if existing:
                existing.is_active = True
                existing.updated_by = request.user.username if request.user.is_authenticated else 'Unknown'
                existing.save(
                    update_fields=['is_active', 'updated_by', 'updated_at'])
                messages.success(request, 'Đã kích hoạt lại hộ khẩu cũ!')
                return redirect('hokhau_list')
            else:
                hokhau = form.save(commit=False)
                hokhau.updated_by = request.user.username if request.user.is_authenticated else 'Unknown'
                hokhau.save()
                messages.success(request, 'Thêm hộ khẩu thành công!')
                return redirect('hokhau_list')
        else:
            # Lấy tất cả lỗi của form (bao gồm cả lỗi trường và lỗi chung)
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(
                        request, f"{form.fields[field].label if field in form.fields else field}: {error}")
            # Nếu có lỗi chung (non_field_errors)
            for error in form.non_field_errors():
                messages.error(request, error)
            return redirect('hokhau_list')
    else:
        form = HoKhauForm(available_canho=available_canho,
                          available_chuho=available_chuho)
    return render(request, 'resident/hokhau_add.html', {'form': form})


@login_required(login_url="login")
@role_required([1])
def hokhau_edit(request, pk):
    hokhau = get_object_or_404(HoKhau, pk=pk)
    old_resident_status = hokhau.resident_status
    # Lấy danh sách căn hộ chưa có hộ khẩu active, cộng thêm căn hộ hiện tại
    canho_with_active_hokhau = HoKhau.objects.filter(
        is_deleted=False, is_active=True, id_canho__isnull=False).values_list('id_canho', flat=True)
    available_canho = CanHo.objects.filter(is_deleted=False).exclude(
        id_canho__in=canho_with_active_hokhau)
    if hokhau.id_canho:
        available_canho = available_canho | CanHo.objects.filter(
            id_canho=hokhau.id_canho.id_canho)
    # Lấy danh sách chủ hộ chưa có hộ khẩu active, cộng thêm chủ hộ hiện tại
    chuho_with_active_hokhau = HoKhau.objects.filter(
        is_deleted=False, is_active=True, id_chuho__isnull=False).values_list('id_chuho', flat=True)
    available_chuho = TaiKhoan.objects.filter(vaitro__id_vaitro=2, is_deleted=False, is_active=True).exclude(
        id_taikhoan__in=chuho_with_active_hokhau)
    if hokhau.id_chuho:
        available_chuho = available_chuho | TaiKhoan.objects.filter(
            id_taikhoan=hokhau.id_chuho.id_taikhoan)

    if request.method == 'POST':
        post = request.POST.copy()
        post['id_chuho'] = str(hokhau.id_chuho.pk)
        
        form = HoKhauForm(post, instance=hokhau, available_canho=available_canho,
                          available_chuho=available_chuho, disabled_chuho=True)
        if form.is_valid():
            hokhau = form.save(commit=False)
            hokhau.updated_by = request.user.username if request.user.is_authenticated else 'Unknown'
            hokhau.save()
            # Nếu resident_status thay đổi thì tạo biến động mới cho tất cả nhân khẩu active thuộc hộ khẩu này
            if old_resident_status != hokhau.resident_status:
                nhankhaus = NhanKhau.objects.filter(
                    id_hokhau=hokhau, is_deleted=False, is_active=True)
                now = timezone.now().date()
                for nk in nhankhaus:
                    BienDongNhanKhau.objects.create(
                        id_nhankhau=nk,
                        loai_biendong=hokhau.resident_status,
                        ngay_batdau=now,
                        ly_do='Cập nhật tình trạng cư trú hộ khẩu',
                        updated_by=request.user.username if request.user.is_authenticated else 'Unknown'
                    )
            messages.success(request, 'Cập nhật hộ khẩu thành công!')
            return redirect('hokhau_list')
        else:
            # Lấy tất cả lỗi của form (bao gồm cả lỗi trường và lỗi chung)
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(
                        request, f"{form.fields[field].label if field in form.fields else field}: {error}")
            # Nếu có lỗi chung (non_field_errors)
            for error in form.non_field_errors():
                messages.error(request, error)
            return redirect('hokhau_list')
    else:
        form = HoKhauForm(instance=hokhau, available_canho=available_canho,
                          available_chuho=available_chuho, disabled_chuho=True)
    return render(request, 'resident/hokhau_edit.html', {'form': form, 'hokhau': hokhau})


@login_required(login_url="login")
@role_required([1])
def hokhau_delete(request, pk):
    hokhau = get_object_or_404(HoKhau, pk=pk)
    if request.method == 'POST':
        # Không cho xóa nếu còn nhân khẩu active
        nhankhau_active_count = NhanKhau.objects.filter(
            id_hokhau=hokhau, is_deleted=False, is_active=True).count()
        if nhankhau_active_count > 0:
            messages.error(
                request, 'Không thể xóa hộ khẩu còn nhân khẩu đang hoạt động!')
            return redirect('hokhau_list')

        hokhau.is_deleted = True
        hokhau.updated_by = request.user.username if request.user.is_authenticated else 'Unknown'
        hokhau.save(update_fields=['is_deleted', 'updated_by', 'updated_at'])
        messages.success(request, 'Đã xóa hộ khẩu!')
        return redirect('hokhau_list')
    return render(request, 'resident/hokhau_delete.html', {'hokhau': hokhau})


@login_required(login_url="login")
@role_required([1, 2])
def hokhau_detail(request, pk):
    hokhau = get_object_or_404(HoKhau, pk=pk)
    nhankhau_list = NhanKhau.objects.filter(
        id_hokhau=hokhau, is_deleted=False, is_active=True)
    nhankhau_count = nhankhau_list.count()
    return render(request, 'resident/hokhau_detail.html', {
        'hokhau': hokhau,
        'nhankhau_list': nhankhau_list,
        'nhankhau_count': nhankhau_count
    })


@login_required(login_url="login")
@role_required([1, 2])
def nhan_khau_profile(request, id_nhankhau):
    nhan_khau = get_object_or_404(NhanKhau, id_nhankhau=id_nhankhau)
    return render(request, "core/nhan_khau_profile.html", {"nhan_khau": nhan_khau})


@login_required(login_url="login")
@role_required([1, 2])
def export_biendong_excel(request):
    if hasattr(request.user, 'vaitro') and getattr(request.user.vaitro, 'id_vaitro', None) == 2:
        biendongs = BienDongNhanKhau.objects.select_related("id_nhankhau").filter(
            id_nhankhau__id_hokhau__id_chuho=request.user
        ).order_by("-ngay_batdau")
    else:
        biendongs = BienDongNhanKhau.objects.select_related(
            "id_nhankhau").order_by("-ngay_batdau")

    wb = Workbook()
    ws = wb.active
    ws.title = "Biến động nhân khẩu"

    header = [
        "ID Biến động",
        "Họ tên",
        "Số căn hộ",
        "CCCD",
        "Quan hệ với chủ hộ",
        "Loại biến động",
        "Ngày bắt đầu",
        "Ngày kết thúc",
        "Lý do",
    ]
    ws.append(header)

    # Style header: fill color, bold, center, wrap text
    header_fill = PatternFill(start_color="FFDEEAF6",
                              end_color="FFDEEAF6", fill_type="solid")
    header_font = Font(bold=True, color="FF222222")
    for col in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True)

    # Set column widths (ID, Họ tên, Số căn hộ, CCCD, Quan hệ, Loại BD, Ngày bắt đầu, Ngày kết thúc, Lý do)
    col_widths = [12, 20, 14, 16, 18, 18, 16, 16, 22]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    for bd in biendongs:
        nk = bd.id_nhankhau
        ws.append([
            bd.id_biendong,
            nk.ho_ten if nk else "",
            nk.id_hokhau.id_canho.so_can_ho if nk and nk.id_hokhau and nk.id_hokhau.id_canho else "",
            nk.cccd if nk else "",
            nk.get_quan_he_chu_ho_display() if nk and nk.quan_he_chu_ho else "",
            bd.get_loai_biendong_display(),
            bd.ngay_batdau.strftime("%d/%m/%Y") if bd.ngay_batdau else "",
            bd.ngay_ketthuc.strftime("%d/%m/%Y") if bd.ngay_ketthuc else "",
            bd.ly_do or "",
        ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="biendong_nhankhau.xlsx"'

    wb.save(response)
    return response


@login_required(login_url="login")
@role_required([1, 2])
def export_nhankhau_excel(request):
    if hasattr(request.user, 'vaitro') and getattr(request.user.vaitro, 'id_vaitro', None) == 2:
        nhankhau = NhanKhau.objects.filter(
            is_deleted=False, is_active=True,
            id_hokhau__is_active=True, id_hokhau__is_deleted=False,
            id_hokhau__id_chuho=request.user
        ).order_by("id_hokhau")
    else:
        nhankhau = NhanKhau.objects.filter(
            is_deleted=False, is_active=True,
            id_hokhau__is_active=True, id_hokhau__is_deleted=False
        ).order_by("id_hokhau")

    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sách nhân khẩu"

    header = [
        "ID Nhân khẩu",
        "Số căn hộ",
        "Họ tên",
        "Giới tính",
        "Ngày sinh",
        "CCCD",
        "Số điện thoại",
        "Email",
        "Quan hệ với chủ hộ",
        "Loại biến động gần nhất",
        "Ngày đăng ký biến động gần nhất",
    ]
    ws.append(header)

    # Style header: fill color, bold, center
    header_fill = PatternFill(start_color="FFDEEAF6",
                              end_color="FFDEEAF6", fill_type="solid")
    header_font = Font(bold=True, color="FF222222")
    for col in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True)

    # Set column widths
    col_widths = [12, 10, 20, 14, 16, 14, 22, 18, 14, 22, 22]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    for nk in nhankhau:
        # Lấy biến động gần nhất
        loai_bd_display = ""
        ngay_bd = ""
        latest_bd = BienDongNhanKhau.objects.filter(
            id_nhankhau=nk, is_deleted=False).order_by('-ngay_batdau').first()

        if latest_bd:
            loai_bd_display = latest_bd.get_loai_biendong_display()
            ngay_bd = latest_bd.ngay_batdau.strftime(
                "%d/%m/%Y") if latest_bd.ngay_batdau else ""

        gioi_tinh_display = nk.get_gioi_tinh_display() if nk.gioi_tinh else ""
        ws.append([
            nk.id_nhankhau,
            nk.id_hokhau.id_canho.so_can_ho if nk.id_hokhau and nk.id_hokhau.id_canho else "",
            nk.ho_ten,
            gioi_tinh_display,
            nk.ngay_sinh,
            nk.cccd,
            nk.so_dien_thoai,
            nk.email,
            nk.get_quan_he_chu_ho_display(),
            loai_bd_display,
            ngay_bd,
        ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="danh_sach_nhan_khau.xlsx"'

    wb.save(response)
    return response


@login_required(login_url="login")
@role_required([1, 2])
def export_hokhau_excel(request):
    if hasattr(request.user, 'vaitro') and getattr(request.user.vaitro, 'id_vaitro', None) == 2:
        hokhau = HoKhau.objects.filter(
            is_deleted=False, is_active=True, id_chuho=request.user)
    else:
        hokhau = HoKhau.objects.filter(is_deleted=False, is_active=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sách hộ khẩu"
    header = [
        "ID Hộ Khẩu",
        "Số căn hộ",
        "Diện tích (m2)",
        "Tài khoản Chủ hộ",
        "Họ tên Chủ hộ",
        "CCCD Chủ hộ",
        "Tình trạng cư trú",
        "Số lượng nhân khẩu",
    ]
    ws.append(header)

    # Style header: fill color, bold, center
    header_fill = PatternFill(start_color="FFDEEAF6",
                              end_color="FFDEEAF6", fill_type="solid")
    header_font = Font(bold=True, color="FF222222")
    for col in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True)

    # Set column widths
    col_widths = [12, 14, 14, 18, 20, 18, 18, 16, 14]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    for hk in hokhau:
        # Lấy chủ hộ từ bảng nhân khẩu (active)
        chu_ho_nk = NhanKhau.objects.filter(
            id_hokhau=hk, is_deleted=False, is_active=True, is_chu_ho=True).first()
        ho_ten_chu_ho = chu_ho_nk.ho_ten if chu_ho_nk else ""
        cccd_chu_ho = chu_ho_nk.cccd if chu_ho_nk else ""

        # Số lượng nhân khẩu active
        so_luong_nk = NhanKhau.objects.filter(
            id_hokhau=hk, is_deleted=False, is_active=True).count()

        # Tình trạng cư trú
        tinh_trang_cu_tru = hk.get_resident_status_display() if hk.resident_status else ""

        ws.append([
            hk.id_hokhau,
            hk.id_canho.so_can_ho if hk.id_canho else "",
            hk.id_canho.dien_tich if hk.id_canho else "",
            hk.id_chuho.username if hk.id_chuho else "",
            ho_ten_chu_ho,
            cccd_chu_ho,
            tinh_trang_cu_tru,
            so_luong_nk
        ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="danh_sach_ho_khau.xlsx"'

    wb.save(response)
    return response


@login_required(login_url="login")
@role_required([1, 2])
def biendong_list(request):
    query = request.GET.get('search', '').strip()
    loai_biendong = request.GET.get('loai_biendong', '').strip()
    date_from = request.GET.get('date_from', '').strip()
    date_to = request.GET.get('date_to', '').strip()

    # Giới hạn vaitro=2 chỉ xem biến động nhân khẩu thuộc hộ khẩu mà mình là chủ hộ
    if hasattr(request.user, 'vaitro') and getattr(request.user.vaitro, 'id_vaitro', None) == 2:
        biendongs = BienDongNhanKhau.objects.filter(
            is_deleted=False,
            id_nhankhau__id_hokhau__id_chuho=request.user
        ).select_related(
            "id_nhankhau", "id_nhankhau__id_hokhau", "id_nhankhau__id_hokhau__id_canho"
        ).order_by('-updated_at')
    else:
        biendongs = BienDongNhanKhau.objects.filter(is_deleted=False).select_related(
            "id_nhankhau", "id_nhankhau__id_hokhau", "id_nhankhau__id_hokhau__id_canho"
        ).order_by('-updated_at')

    if query:
        biendongs = biendongs.filter(
            Q(id_nhankhau__ho_ten__icontains=query) |
            Q(id_nhankhau__cccd__icontains=query) |
            Q(id_nhankhau__id_hokhau__id_canho__so_can_ho__icontains=query)
        )
    if loai_biendong:
        biendongs = biendongs.filter(loai_biendong=loai_biendong)
    if date_from:
        biendongs = biendongs.filter(ngay_batdau__gte=date_from)
    if date_to:
        biendongs = biendongs.filter(ngay_batdau__lte=date_to)
    biendongs = biendongs.order_by('-updated_at')

    return render(
        request,
        "resident/biendongnhankhau/biendong_list.html",
        {
            "biendongs": biendongs,
            "query": query,
            "loai_biendong": loai_biendong,
            "date_from": date_from,
            "date_to": date_to,
            "loai_bd_choices": LoaiBienDong.choices,
        },
    )


@login_required(login_url="login")
@role_required([1, 2])
def dang_ky_bdnk(request, id_nhankhau):
    nhan_khau = get_object_or_404(NhanKhau, id_nhankhau=id_nhankhau)
    # Phân quyền: vaitro=2 chỉ được đăng ký cho nhân khẩu thuộc hộ khẩu mình
    if hasattr(request.user, 'vaitro') and getattr(request.user.vaitro, 'id_vaitro', None) == 2:
        if not nhan_khau.id_hokhau or nhan_khau.id_hokhau.id_chuho != request.user:
            messages.error(
                request, 'Bạn không có quyền đăng ký biến động cho nhân khẩu này!')
            return redirect('nhankhau_list')

    if request.method == "POST":
        form = BienDongForm(request.POST, nhan_khau=nhan_khau)
        if form.is_valid():
            bien_dong = form.save(commit=False)
            bien_dong.id_nhankhau = nhan_khau
            bien_dong.updated_by = request.user.username if request.user.is_authenticated else 'Unknown'
            bien_dong.save()
            # Cập nhật ngày kết thúc cho bản ghi chưa xóa liền trước nếu đang null
            prev_bd = nhan_khau.nhan_khau.filter(is_deleted=False).exclude(
                id_biendong=bien_dong.id_biendong).order_by('-ngay_batdau').first()
            if prev_bd and prev_bd.ngay_ketthuc is None and prev_bd.ngay_batdau < bien_dong.ngay_batdau:
                prev_bd.ngay_ketthuc = bien_dong.ngay_batdau
                prev_bd.updated_by = request.user.username if request.user.is_authenticated else 'Unknown'
                prev_bd.save(update_fields=['ngay_ketthuc', 'updated_by'])
            # Nếu loại biến động là chấm dứt cư trú thì set is_active=False cho nhân khẩu
            if bien_dong.loai_biendong == 'cham_dut':
                nhan_khau.is_active = False
                nhan_khau.updated_by = request.user.username if request.user.is_authenticated else 'Unknown'
                nhan_khau.save(update_fields=['is_active', 'updated_by'])
            messages.success(
                request, "Đăng ký biến động nhân khẩu thành công!")
            return redirect("nhankhau_list")
        else:
            # Lấy tất cả lỗi của form (bao gồm cả lỗi trường và lỗi chung)
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(
                        request, f"{form.fields[field].label if field in form.fields else field}: {error}")
            # Nếu có lỗi chung (non_field_errors)
            for error in form.non_field_errors():
                messages.error(request, error)
            return redirect('nhankhau_list')
    else:
        form = BienDongForm(nhan_khau=nhan_khau)
    return render(
        request,
        "resident/biendongnhankhau/dangkybdnk.html",
        {"nhan_khau": nhan_khau, "form": form},
    )


@login_required(login_url="login")
@role_required([1, 2])
def biendong_delete(request, id_biendong):
    # Nếu là cư dân chỉ cho xóa biến động của nhân khẩu thuộc hộ khẩu mình
    if hasattr(request.user, 'vaitro') and getattr(request.user.vaitro, 'id_vaitro', None) == 2:
        biendong = get_object_or_404(
            BienDongNhanKhau,
            pk=id_biendong,
            is_deleted=False,
            id_nhankhau__id_hokhau__id_chuho=request.user
        )
    else:
        biendong = get_object_or_404(
            BienDongNhanKhau, pk=id_biendong, is_deleted=False)
    if request.method == 'POST':
        # Nếu biến động là chấm dứt cư trú thì cập nhật lại trạng thái nhân khẩu
        if biendong.loai_biendong == 'cham_dut' and biendong.id_nhankhau:
            nhankhau = biendong.id_nhankhau
            nhankhau.is_active = True
            nhankhau.updated_by = request.user.username if request.user.is_authenticated else 'Unknown'
            nhankhau.save(update_fields=['is_active', 'updated_by'])
        
        biendong.is_deleted = True
        biendong.updated_by = request.user.username if request.user.is_authenticated else 'Unknown'
        biendong.save(update_fields=['is_deleted', 'updated_by', 'updated_at'])
        messages.success(request, 'Đã xóa biến động nhân khẩu!')
        return redirect('biendong_list')
    return render(request, 'resident/biendongnhankhau/confirm_delete.html', {'biendong': biendong})
